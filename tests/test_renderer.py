"""Generic renderer smoke tests; no personal resume data."""

import importlib.util
import tempfile
import unittest
from pathlib import Path

from docx import Document
from pypdf import PdfReader


SCRIPT = Path(__file__).resolve().parents[1] / "scripts" / "render_resume.py"
SPEC = importlib.util.spec_from_file_location("render_resume", SCRIPT)
renderer = importlib.util.module_from_spec(SPEC)
SPEC.loader.exec_module(renderer)


class RendererTest(unittest.TestCase):
    def test_four_formats_and_narrative_guards(self):
        with tempfile.TemporaryDirectory() as directory:
            root = Path(directory)
            data = root / "data"
            data.mkdir()
            (data / "previous-resume.txt").write_text(
                "The source had a long distinctive sentence about a fictional service and its release checks.\n",
                encoding="utf-8",
            )
            (data / "~$lock.docx").write_bytes(b"office lock file, not a document")
            source = root / "resume.md"
            source.write_text(
                "# Sample Candidate\nCity, ST | sample@example.invalid\n"
                "## Summary\nSoftware test engineer focused on reliable releases.\n"
                "## Experience\n### Test Engineer | Example Company | 2020 to 2024\n"
                "- Built regression checks for a service. Improved release confidence for the team.\n"
                "## Skills\nPython, API testing, automation\n",
                encoding="utf-8",
            )
            blocks = renderer.parse_markdown(source.read_text(encoding="utf-8"))
            renderer.validate(blocks, data)
            renderer.write_docx(blocks, root / "resume.docx")
            renderer.write_pdf(blocks, root / "resume.pdf")
            renderer.write_txt(blocks, root / "resume.txt")
            phrase = "Improved release confidence for the team"
            self.assertIn(phrase, " ".join(p.text for p in Document(root / "resume.docx").paragraphs))
            self.assertIn(phrase, PdfReader(root / "resume.pdf").pages[0].extract_text())
            self.assertIn("• Built regression checks", (root / "resume.txt").read_text(encoding="utf-8"))
            headings = [p for p in Document(root / "resume.docx").paragraphs if p.text in {"SUMMARY", "EXPERIENCE", "SKILLS"}]
            self.assertEqual(len(headings), 3)
            self.assertTrue(all(p.alignment == 1 for p in headings))
            self.assertTrue(source.is_file())
            bad_first_person = renderer.parse_markdown("# Sample Candidate\n## Experience\n- I built a framework.\n")
            with self.assertRaises(ValueError):
                renderer.validate(bad_first_person, data)
            renderer.validate(renderer.parse_markdown("# Sample Candidate\n## Experience\n- Built end-to-end checks.\n"), data)
            for dashed in ("- Built checks - quickly.", "- Built checks — quickly.", "- Built checks -- quickly."):
                with self.assertRaises(ValueError):
                    renderer.validate(renderer.parse_markdown(f"# Sample Candidate\n## Experience\n{dashed}\n"), data)
            vocabulary = root / "vocabulary.md"
            vocabulary.write_text(
                "| Write | Meaning | Avoid |\n|---|---|---|\n| test environment | Shared test setup | sandbox, Sandbox |\n",
                encoding="utf-8",
            )
            renderer.validate(blocks, data, vocabulary)
            with self.assertRaises(ValueError):
                renderer.validate(renderer.parse_markdown("# Sample Candidate\n## Experience\n- Built a sandbox.\n"), data, vocabulary)
            bad_copy = renderer.parse_markdown(
                "# Sample Candidate\n## Experience\n"
                "- The source had a long distinctive sentence about a fictional service and its release checks.\n"
            )
            with self.assertRaises(ValueError):
                renderer.validate(bad_copy, data)


class ContactFileTest(unittest.TestCase):
    BODY = "## Sample Role\n\nBuilds reliable checks.\n\n## Experience\n\n### Engineer | Example Corp | 2020 to 2024\n\n- Built end-to-end checks for payment flows.\n"
    CONTACT = "# Pat Sampleton\nSampleville, ST | pat@example.invalid | 555.010.0199\n"

    def setUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.addCleanup(self.directory.cleanup)
        self.root = Path(self.directory.name)
        (self.root / "data").mkdir()
        (self.root / "data" / "contact.md").write_text(self.CONTACT, encoding="utf-8")
        self.body = self.root / "generic.md"
        self.body.write_text(self.BODY, encoding="utf-8")

    def test_body_is_joined_to_contact_and_rendered_under_output(self):
        folder = renderer.render(self.body, self.root)
        self.assertEqual(folder, (self.root / "output" / "generic").resolve())
        for name in ("resume.md", "resume.docx", "resume.pdf", "resume.txt"):
            self.assertTrue((folder / name).is_file(), name)
        self.assertTrue((folder / "resume.md").read_text(encoding="utf-8").startswith("# Pat Sampleton\n"))
        self.assertIn("Pat Sampleton", PdfReader(folder / "resume.pdf").pages[0].extract_text())
        self.assertNotIn("Pat Sampleton", self.body.read_text(encoding="utf-8"))
        self.assertEqual(list(self.root.glob("generic.*")), [self.body])

    def test_output_outside_output_dir_is_refused(self):
        with self.assertRaises(ValueError):
            renderer.render(self.body, self.root, out_dir=self.root / "elsewhere")

    def test_body_repeating_contact_details_is_refused(self):
        for leak in ("Reach me at pat@example.invalid", "Call 555.010.0199", "Sampleton led it", "Based in Sampleville, ST"):
            self.body.write_text(self.BODY + f"\n{leak}.\n", encoding="utf-8")
            with self.assertRaises(ValueError, msg=leak):
                renderer.render(self.body, self.root)

    def test_missing_or_malformed_contact_file_is_refused(self):
        (self.root / "data" / "contact.md").unlink()
        with self.assertRaises(ValueError):
            renderer.render(self.body, self.root)
        (self.root / "data" / "contact.md").write_text("# Pat Sampleton\n", encoding="utf-8")
        with self.assertRaises(ValueError):
            renderer.render(self.body, self.root)

    def test_headed_resume_keeps_working_and_rejects_contact_options(self):
        headed = self.root / "resume.md"
        headed.write_text(self.CONTACT.replace("\n", "\n\n", 1) + self.BODY, encoding="utf-8")
        self.assertEqual(renderer.render(headed, self.root), self.root)
        self.assertTrue((self.root / "resume.pdf").is_file())
        with self.assertRaises(ValueError):
            renderer.render(headed, self.root, contact_path=self.root / "data" / "contact.md")


if __name__ == "__main__":
    unittest.main()
