"""Validate a resume and write resume.docx, resume.pdf, and resume.txt.

Run from the repository root:
  python scripts/render_resume.py output/<version>/resume.md
  python scripts/render_resume.py <body>.md [--contact data/contact.md] [--out output/<name>]

A resume.md that starts with "# Name" renders beside itself. A body file with no header (it
starts at "## ...") is joined to the private contact file (two lines: "# Name" and a contact
line) and rendered under output/ only, so a file holding contact details never lands in a
tracked location.
"""

from __future__ import annotations

import argparse
import re
from pathlib import Path
from xml.sax.saxutils import escape

from docx import Document
from docx.shared import Inches, Pt
from pypdf import PdfReader
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import inch
from reportlab.platypus import Paragraph, SimpleDocTemplate


Block = tuple[str, str]
WORD_RE = re.compile(r"[a-z0-9]+(?:'[a-z]+)?", re.I)


def words(text: str) -> list[str]:
    return [match.lower() for match in WORD_RE.findall(text)]


def parse_markdown(text: str) -> list[Block]:
    blocks: list[Block] = []
    for raw in text.splitlines():
        line = raw.strip()
        if not line:
            continue
        for prefix, kind in (("### ", "role"), ("## ", "section"), ("# ", "name"), ("- ", "bullet")):
            if line.startswith(prefix):
                value = line[len(prefix):].strip()
                break
        else:
            kind, value = "text", line
        if not value or any(token in value for token in ("**", "[", "](", "|---")):
            raise ValueError("Use nonempty plain text without inline Markdown or tables")
        blocks.append((kind, value))
    if not blocks or blocks[0][0] != "name":
        raise ValueError("Resume must begin with '# Candidate Name'")
    return blocks


def extract_source_text(path: Path) -> str:
    suffix = path.suffix.lower()
    if suffix in {".md", ".txt"}:
        return path.read_text(encoding="utf-8", errors="replace")
    if suffix == ".docx":
        return "\n".join(p.text for p in Document(path).paragraphs)
    if suffix == ".pdf":
        return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)
    return ""


def copied_prose(blocks: list[Block], data_dir: Path, width: int = 10) -> str | None:
    """Detect long verbatim spans in prose; headings and names are excluded."""
    if not data_dir.is_dir():
        return None
    candidates = [words(value) for kind, value in blocks if kind in {"bullet", "text"}]
    if candidates and blocks[1:2] and blocks[1][0] == "text":
        candidates.pop(0)  # contact line
    candidate_grams = {
        tuple(tokens[i:i + width])
        for tokens in candidates
        for i in range(max(0, len(tokens) - width + 1))
    }
    if not candidate_grams:
        return None
    for path in sorted(data_dir.rglob("*")):
        if (not path.is_file() or path.name.startswith("~$")
                or path.suffix.lower() not in {".md", ".txt", ".docx", ".pdf"}):
            continue
        try:
            source = words(extract_source_text(path))
        except Exception as exc:
            raise ValueError(f"Could not read source file {path.name}: {exc}") from exc
        if any(tuple(source[i:i + width]) in candidate_grams for i in range(max(0, len(source) - width + 1))):
            return path.name
    return None


# Dashes used as punctuation: en/em dashes and similar, a spaced hyphen, or a double hyphen.
# A hyphen inside a compound word (end-to-end) is allowed.
DASH_RE = re.compile(r"[‒-―−]|\s-|-\s|--")


def load_vocabulary(path: Path | None) -> list[tuple[str, str]]:
    """Return (avoid, write) pairs from the Write | Meaning | Avoid tables in vocabulary.md."""
    if path is None or not path.is_file():
        return []
    pairs = []
    for line in path.read_text(encoding="utf-8").splitlines():
        cols = [col.strip() for col in line.strip().strip("|").split("|")]
        if len(cols) < 3 or cols[0] == "Write" or not cols[0].strip("-: "):
            continue
        pairs.extend((avoid.strip(), cols[0]) for avoid in cols[2].split(",") if avoid.strip())
    return pairs


def validate(blocks: list[Block], data_dir: Path, vocabulary_path: Path | None = None) -> None:
    vocabulary = load_vocabulary(vocabulary_path)
    for index, (kind, value) in enumerate(blocks):
        # The name and contact line may legitimately contain hyphens (emails, URLs).
        if kind == "name" or (index == 1 and kind == "text"):
            continue
        if DASH_RE.search(value):
            raise ValueError("Dash used as sentence punctuation; split into separate sentences")
        for avoid, write in vocabulary:
            if re.search(rf"(?<!\w){re.escape(avoid)}(?!\w)", value):
                raise ValueError(f"Use '{write}' instead of '{avoid}' (see vocabulary.md)")
        if kind not in {"bullet", "text"} or (index == 1 and kind == "text"):
            continue
        if re.search(r"\b(?:I|me|my|we|our)\b", value, re.I):
            raise ValueError("First-person or conversational wording in resume prose")
        tokens = words(value)
        if len(tokens) > (45 if kind == "bullet" else 50):
            raise ValueError("Resume bullet or paragraph is too long; condense it")
        for sentence in re.split(r"(?<=[.!?])\s+", value):
            if len(words(sentence)) > 27:
                raise ValueError("Run-on sentence detected; use shorter sentences")
    copied_from = copied_prose(blocks, data_dir)
    if copied_from:
        raise ValueError(f"Resume repeats a long source phrase from {copied_from}; rewrite it")


def write_docx(blocks: list[Block], target: Path) -> None:
    doc = Document()
    page = doc.sections[0]
    page.top_margin = page.bottom_margin = Inches(0.5)
    page.left_margin = page.right_margin = Inches(0.6)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(2)
    normal.paragraph_format.line_spacing = 1.0
    for index, (kind, value) in enumerate(blocks):
        if kind == "name":
            para = doc.add_paragraph()
            para.alignment = 1
            run = para.add_run(value)
            run.bold = True
            run.font.size = Pt(16)
        elif kind == "section":
            para = doc.add_paragraph()
            para.alignment = 1
            para.paragraph_format.space_before = Pt(6)
            para.paragraph_format.keep_with_next = True
            para.add_run(value.upper()).bold = True
        elif kind == "role":
            para = doc.add_paragraph()
            para.paragraph_format.space_before = Pt(3)
            para.paragraph_format.keep_with_next = True
            para.add_run(value).bold = True
        elif kind == "bullet":
            para = doc.add_paragraph(style="List Bullet")
            para.paragraph_format.left_indent = Inches(0.22)
            para.paragraph_format.first_line_indent = Inches(-0.12)
            para.add_run(value)
        else:
            para = doc.add_paragraph(value)
            if index == 1:
                para.alignment = 1
    doc.save(target)


def write_pdf(blocks: list[Block], target: Path) -> None:
    styles = {
        "name": ParagraphStyle("name", fontName="Helvetica-Bold", fontSize=16, leading=18, alignment=TA_CENTER, spaceAfter=1),
        "contact": ParagraphStyle("contact", fontName="Helvetica", fontSize=9, leading=11, alignment=TA_CENTER, spaceAfter=3),
        "section": ParagraphStyle("section", fontName="Helvetica-Bold", fontSize=10, leading=12, textColor=colors.HexColor("#18354D"), alignment=TA_CENTER, spaceBefore=6, spaceAfter=2, keepWithNext=True),
        "role": ParagraphStyle("role", fontName="Helvetica-Bold", fontSize=10, leading=12, spaceBefore=3, spaceAfter=1, keepWithNext=True),
        "text": ParagraphStyle("text", fontName="Helvetica", fontSize=10, leading=12, spaceAfter=2),
        "bullet": ParagraphStyle("bullet", fontName="Helvetica", fontSize=10, leading=12, leftIndent=15, bulletIndent=5, firstLineIndent=0, spaceAfter=2),
    }
    story = []
    for index, (kind, value) in enumerate(blocks):
        style = "contact" if kind == "text" and index == 1 else kind
        text = value.upper() if kind == "section" else value
        story.append(Paragraph(escape(text), styles[style], bulletText="•" if kind == "bullet" else None))
    document = SimpleDocTemplate(str(target), pagesize=letter, leftMargin=0.6 * inch,
                                 rightMargin=0.6 * inch, topMargin=0.5 * inch,
                                 bottomMargin=0.5 * inch, title=blocks[0][1] + " Resume")
    document.build(story)


def write_txt(blocks: list[Block], target: Path) -> None:
    lines: list[str] = []
    for kind, value in blocks:
        if kind == "section":
            lines.extend(("", value.upper()))
        elif kind == "role":
            lines.extend(("", value))
        elif kind == "bullet":
            lines.append("• " + value)
        else:
            lines.append(value)
    target.write_text("\n".join(lines).strip() + "\n", encoding="utf-8")


EMAIL_RE = re.compile(r"[\w.+-]+@[\w-]+(?:\.[\w-]+)+")
PHONE_RE = re.compile(r"(?<!\d)(?:\+?1[\s.-]?)?\(?\d{3}\)?[\s.-]?\d{3}[\s.-]?\d{4}(?!\d)")


def is_headerless(text: str) -> bool:
    return not text.lstrip().startswith("# ")


def check_no_identifiers(body: str, contact: str) -> None:
    """Refuse a body that repeats the name or contact details from the contact file."""
    if EMAIL_RE.search(body) or PHONE_RE.search(body):
        raise ValueError("Resume body contains an email address or phone number; keep it in the contact file")
    lowered = body.lower()
    name_line, _, contact_line = contact.strip().partition("\n")
    words_in_name = [t for t in re.split(r"[\s,]+", name_line.lstrip("# ")) if len(t) >= 4]
    fields = [f.strip() for f in contact_line.split("|") if len(f.strip()) >= 4]
    for term in words_in_name + fields:
        if re.search(rf"(?<!\w){re.escape(term.lower())}(?!\w)", lowered):
            raise ValueError("Resume body repeats a name or contact detail from the contact file")


def load_contact(path: Path) -> str:
    if not path.is_file():
        raise ValueError(f"Contact file not found: {path} (two lines: '# Name', then a contact line)")
    lines = [line for line in path.read_text(encoding="utf-8").splitlines() if line.strip()]
    if len(lines) != 2 or not lines[0].startswith("# "):
        raise ValueError("Contact file must have exactly two lines: '# Name' and a contact line")
    return "\n".join(lines)


def render(source: Path, repo_root: Path, contact_path: Path | None = None, out_dir: Path | None = None) -> Path:
    """Validate and render one resume; return the directory the files were written to."""
    text = source.read_text(encoding="utf-8")
    if is_headerless(text):
        output_root = (repo_root / "output").resolve()
        out_dir = (out_dir or output_root / source.stem).resolve()
        if out_dir != output_root and output_root not in out_dir.parents:
            raise ValueError("A headerless body renders only under output/ (it gains contact details)")
        contact = load_contact(contact_path or repo_root / "data" / "contact.md")
        check_no_identifiers(text, contact)
        out_dir.mkdir(parents=True, exist_ok=True)
        target = out_dir / "resume.md"
        target.write_text(contact + "\n\n" + text.lstrip(), encoding="utf-8")
    else:
        if source.name != "resume.md" or out_dir is not None or contact_path is not None:
            raise ValueError("A resume with its own header must be named resume.md; --out and --contact do not apply")
        target = source
    blocks = parse_markdown(target.read_text(encoding="utf-8"))
    validate(blocks, repo_root / "data", repo_root / "vocabulary.md")
    write_docx(blocks, target.with_suffix(".docx"))
    write_pdf(blocks, target.with_suffix(".pdf"))
    write_txt(blocks, target.with_suffix(".txt"))
    return target.parent


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    parser.add_argument("source", type=Path, help="resume.md with a header, or a headerless body .md")
    parser.add_argument("--contact", type=Path, help="contact file for a headerless body (default data/contact.md)")
    parser.add_argument("--out", type=Path, help="output folder under output/ for a headerless body")
    args = parser.parse_args()
    if not args.source.is_file():
        parser.error("source file not found")
    repo_root = Path(__file__).resolve().parent.parent
    try:
        folder = render(args.source, repo_root, args.contact, args.out)
    except ValueError as exc:
        parser.error(str(exc))
    print("Rendered DOCX, PDF, TXT, and Markdown in", folder)


if __name__ == "__main__":
    main()
